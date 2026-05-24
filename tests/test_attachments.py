import pytest

from jarvis_app.core.attachments import (
    AttachmentError,
    AttachedFile,
    attachment_preview,
    build_attachment_context,
    build_user_message_with_attachments,
    format_file_size,
    read_attachment,
)


def test_read_attachment_accepts_supported_text_file(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# Notes\nLocal Jarvis context.", encoding="utf-8")

    attached = read_attachment(path)

    assert attached.name == "notes.md"
    assert attached.extension == ".md"
    assert attached.size_bytes > 0
    assert "Local Jarvis context" in attached.content
    assert not attached.truncated


def test_read_attachment_rejects_unsupported_file_type(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"not really an image")

    with pytest.raises(AttachmentError):
        read_attachment(path)


def test_read_attachment_extracts_docx_text(tmp_path):
    from docx import Document

    path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Jarvis DOCX extraction works.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Jarvis"
    document.save(path)

    attached = read_attachment(path)

    assert attached.extension == ".docx"
    assert "Jarvis DOCX extraction works." in attached.content
    assert "Name | Jarvis" in attached.content


def test_read_attachment_extracts_pdf_text(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    path = tmp_path / "sample.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 40 240 Td (Jarvis PDF extraction works.) Tj ET")
    page[NameObject("/Contents")] = stream
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    with path.open("wb") as file:
        writer.write(file)

    attached = read_attachment(path)

    assert attached.extension == ".pdf"
    assert "Jarvis PDF extraction works" in attached.content


def test_build_attachment_context_includes_file_names(tmp_path):
    path = tmp_path / "code.py"
    path.write_text("print('hello')", encoding="utf-8")
    attached = read_attachment(path)

    context = build_attachment_context([attached])

    assert "code.py" in context
    assert "TYPE: .py" in context
    assert "SIZE:" in context
    assert "print('hello')" in context


def test_build_user_message_with_attachments_places_request_after_context(tmp_path):
    path = tmp_path / "facts.txt"
    path.write_text("Project codename: Blue Lantern", encoding="utf-8")
    attached = read_attachment(path)

    message = build_user_message_with_attachments("What is the codename?", [attached])

    assert "Project codename: Blue Lantern" in message
    assert message.endswith("USER REQUEST:\nWhat is the codename?")


def test_attachment_preview_compacts_whitespace(tmp_path):
    path = tmp_path / "preview.txt"
    path.write_text("line one\n\nline two", encoding="utf-8")
    attached = read_attachment(path)

    assert attachment_preview(attached) == "line one line two"


def test_format_file_size_uses_readable_units():
    assert format_file_size(512) == "512 B"
    assert format_file_size(2048) == "2.0 KB"


def test_build_attachment_context_marks_truncated_document(tmp_path):
    attached = AttachedFile(
        path=tmp_path / "big.pdf",
        name="big.pdf",
        extension=".pdf",
        size_bytes=525_000,
        content="A" * 100,
        truncated=True,
    )

    context = build_attachment_context([attached])

    assert "TYPE: .pdf" in context
    assert "Content truncated for V1 context limit" in context
