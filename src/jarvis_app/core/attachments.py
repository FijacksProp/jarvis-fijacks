from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


SUPPORTED_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".py",
    ".js",
    ".ts",
    ".json",
    ".html",
    ".css",
    ".csv",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".log",
}
SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
}
SUPPORTED_EXTENSIONS = SUPPORTED_TEXT_EXTENSIONS | SUPPORTED_DOCUMENT_EXTENSIONS
MAX_FILE_BYTES = 5_000_000
MAX_TEXT_CHARS_PER_FILE = 12_000
MAX_DOCUMENT_CHARS_PER_FILE = 3_500


class AttachmentError(RuntimeError):
    """Raised when a selected attachment cannot be read safely."""


@dataclass(frozen=True, slots=True)
class AttachedFile:
    path: Path
    name: str
    extension: str
    size_bytes: int
    content: str
    truncated: bool = False


def read_attachment(path: str | Path) -> AttachedFile:
    file_path = Path(path).expanduser().resolve()
    if not file_path.exists() or not file_path.is_file():
        raise AttachmentError(f"File does not exist: {file_path}")

    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise AttachmentError(f"Unsupported file type: {file_path.suffix or '(none)'}. Use: {supported}")

    size = file_path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise AttachmentError(
            f"File is too large for V1 attachments: {file_path.name} "
            f"({size / 1_000_000:.2f} MB). Limit is 5 MB."
        )

    content = _extract_content(file_path, extension)
    char_limit = _char_limit_for_extension(extension)
    truncated = len(content) > char_limit
    if truncated:
        content = content[:char_limit]

    return AttachedFile(
        path=file_path,
        name=file_path.name,
        extension=extension or "(none)",
        size_bytes=size,
        content=content,
        truncated=truncated,
    )


def build_attachment_context(files: list[AttachedFile]) -> str:
    if not files:
        return ""

    sections = [
        "The user attached these local files. Read the file content before answering. "
        "Answer from the attached files first. If the answer is not present in the files, say so clearly. "
        "If a file is truncated, answer from the visible excerpt and mention that the full document was not included.",
    ]
    for file in files:
        truncation_note = "\n[Content truncated for V1 context limit.]" if file.truncated else ""
        sections.append(
            f"\n--- FILE: {file.name}\n"
            f"PATH: {file.path}\n"
            f"TYPE: {file.extension}\n"
            f"SIZE: {format_file_size(file.size_bytes)}\n"
            f"CONTENT:\n{file.content}{truncation_note}\n---"
        )
    return "\n".join(sections)


def build_user_message_with_attachments(message: str, files: list[AttachedFile]) -> str:
    context = build_attachment_context(files)
    if not context:
        return message
    return f"{context}\n\nUSER REQUEST:\n{message}"


def attachment_preview(file: AttachedFile, max_chars: int = 420) -> str:
    content = " ".join(file.content.split())
    if not content:
        return "(empty file)"
    if len(content) <= max_chars:
        return content
    return f"{content[:max_chars].rstrip()}..."


def format_file_size(size_bytes: int) -> str:
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    return f"{size_bytes / (1024 * 1024):.2f} MB"


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise AttachmentError(f"Could not read file as text: {path}")


def _extract_content(path: Path, extension: str) -> str:
    if extension == ".pdf":
        return _read_pdf(path)
    if extension == ".docx":
        return _read_docx(path)
    return _read_text(path)


def _char_limit_for_extension(extension: str) -> int:
    if extension in SUPPORTED_DOCUMENT_EXTENSIONS:
        return MAX_DOCUMENT_CHARS_PER_FILE
    return MAX_TEXT_CHARS_PER_FILE


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise AttachmentError("PDF support requires pypdf. Install project dependencies again.") from exc

    try:
        reader = PdfReader(str(path))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {index}]\n{text.strip()}")
        content = "\n\n".join(pages).strip()
    except Exception as exc:
        raise AttachmentError(f"Could not extract text from PDF: {path.name}") from exc

    if not content:
        raise AttachmentError(
            f"No readable text found in PDF: {path.name}. Scanned/image PDFs need OCR, which is not in V1 yet."
        )
    return content


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise AttachmentError("DOCX support requires python-docx. Install project dependencies again.") from exc

    try:
        document = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_rows = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_rows.append(" | ".join(cells))
        content = "\n".join(paragraphs + table_rows).strip()
    except Exception as exc:
        raise AttachmentError(f"Could not extract text from DOCX: {path.name}") from exc

    if not content:
        raise AttachmentError(f"No readable text found in DOCX: {path.name}")
    return content
